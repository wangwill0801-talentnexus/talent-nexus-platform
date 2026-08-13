from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(r"C:\Users\William Wang\Desktop\TalentNexus-104-Golden")
OUT = ROOT / "docs" / "Talent Nexus 英鏈人才_Connector完整操作手冊_v1.0_正式版.docx"
LOGO = Path(r"C:\Users\William Wang\Pictures\logo2.jpg")

BLUE = "2E74B5"
DEEP_BLUE = "1F4D78"
LIGHT_BLUE = "E8EEF5"
PALE_BLUE = "F4F8FC"
TEXT = "243447"
GREY = "5F6B76"
TABLE_WIDTH = 9360


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge in kwargs:
            edge_data = kwargs.get(edge)
            tag = "w:%s" % edge
            element = borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                borders.append(element)
            for key in ["val", "sz", "space", "color"]:
                if key in edge_data:
                    element.set(qn("w:%s" % key), str(edge_data[key]))


def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH))
    tbl_w.set(qn("w:type"), "dxa")
    ind = tbl_pr.find(qn("w:tblInd"))
    if ind is None:
        ind = OxmlElement("w:tblInd")
        tbl_pr.append(ind)
    ind.set(qn("w:w"), "0")
    ind.set(qn("w:type"), "dxa")
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_mar = cell._tc.get_or_add_tcPr().find(qn("w:tcMar"))
            if tc_mar is None:
                tc_mar = OxmlElement("w:tcMar")
                cell._tc.get_or_add_tcPr().append(tc_mar)
            for side, val in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
                node = tc_mar.find(qn("w:%s" % side))
                if node is None:
                    node = OxmlElement("w:%s" % side)
                    tc_mar.append(node)
                node.set(qn("w:w"), str(val))
                node.set(qn("w:type"), "dxa")


def set_run_font(run, size=11, color=TEXT, bold=False, italic=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def style_paragraph(p, after=6, before=0, line=1.25):
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing = line


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    style_paragraph(p)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    style_paragraph(p, after=4)
    p.paragraph_format.left_indent = Inches(0.375 + 0.2 * level)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    style_paragraph(p, after=4)
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style="Heading %d" % level)
    r = p.add_run(text)
    if level == 1:
        set_run_font(r, 16, BLUE, bold=True)
    elif level == 2:
        set_run_font(r, 13, BLUE, bold=True)
    else:
        set_run_font(r, 12, DEEP_BLUE, bold=True)
    return p


def add_callout(doc, title, body, fill=PALE_BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [TABLE_WIDTH])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, top={"val": "single", "sz": 6, "color": BLUE}, bottom={"val": "single", "sz": 6, "color": BLUE}, left={"val": "single", "sz": 6, "color": BLUE}, right={"val": "single", "sz": 6, "color": BLUE})
    p = cell.paragraphs[0]
    style_paragraph(p, after=2)
    r = p.add_run(title)
    set_run_font(r, 11, DEEP_BLUE, bold=True)
    p2 = cell.add_paragraph()
    style_paragraph(p2, after=0)
    r2 = p2.add_run(body)
    set_run_font(r2)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, LIGHT_BLUE)
        set_cell_border(cell, top={"val": "single", "sz": 6, "color": "B7C9DA"}, bottom={"val": "single", "sz": 6, "color": "B7C9DA"}, left={"val": "single", "sz": 4, "color": "D7E1EA"}, right={"val": "single", "sz": 4, "color": "D7E1EA"})
        p = cell.paragraphs[0]
        style_paragraph(p, after=0)
        r = p.add_run(h)
        set_run_font(r, bold=True, color=DEEP_BLUE)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_border(cells[i], top={"val": "single", "sz": 4, "color": "D7E1EA"}, bottom={"val": "single", "sz": 4, "color": "D7E1EA"}, left={"val": "single", "sz": 4, "color": "D7E1EA"}, right={"val": "single", "sz": 4, "color": "D7E1EA"})
            p = cells[i].paragraphs[0]
            style_paragraph(p, after=0)
            r = p.add_run(value)
            set_run_font(r)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def configure_doc(doc):
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.492)
    sec.footer_distance = Inches(0.492)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in (("Heading 1", 16, BLUE, 18, 10), ("Heading 2", 13, BLUE, 14, 7), ("Heading 3", 12, DEEP_BLUE, 10, 5)):
        st = doc.styles[name]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.line_spacing = 1.15
    for list_name in ("List Bullet", "List Number"):
        st = doc.styles[list_name]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
        st.font.size = Pt(11)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.line_spacing = 1.25
    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("Talent Nexus Connector  ·  內部操作手冊")
    set_run_font(r, 8, GREY)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("Talent Nexus 英鏈人才  ·  v1.0  ·  正式版")
    set_run_font(r, 8, GREY)


def build():
    doc = Document()
    configure_doc(doc)

    # Editorial cover
    if LOGO.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(18)
        p.add_run().add_picture(str(LOGO), width=Inches(2.15))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, after=8)
    r = p.add_run("TALENT NEXUS CONNECTOR")
    set_run_font(r, 10, BLUE, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, after=7)
    r = p.add_run("完整操作手冊")
    set_run_font(r, 27, DEEP_BLUE, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, after=20)
    r = p.add_run("ATS 履歷解析、AI Fill、AI 優化與語言轉換")
    set_run_font(r, 14, GREY)
    meta = add_table(doc, ["文件資訊", "內容"], [
        ("文件編號", "TN-CON-OPS-001"),
        ("版本／狀態", "v1.0／正式版"),
        ("適用版本", "Talent Nexus Connector 5000.0.113"),
        ("適用對象", "招聘顧問"),
        ("發布日期", "2026-08-10"),
    ], [2700, 6660])
    add_callout(doc, "本手冊的核心原則", "Connector 只協助擷取、整理及填入目前 ATS 表單；最後的內容確認與 ATS Save 永遠由招聘顧問執行。未在履歷中出現的資料保持空白，不以 AI 推測。", fill="EEF6F8")
    doc.add_page_break()

    add_heading(doc, "1. 目的與適用範圍")
    add_body(doc, "本手冊說明 Talent Nexus Connector 在品聘中文 ATS、英文 ATS、104、LinkedIn、PDF 及 DOCX 履歷來源上的日常操作方式，並特別涵蓋新增人選與既有人選編輯頁面的 AI Fill 流程。")
    add_body(doc, "本版本是招聘顧問的操作指南，不取代 ATS 原有的權限、分類選項或 Save 流程。")

    add_heading(doc, "2. 功能總覽")
    for t in [
        "從 PDF 或 DOCX 先在本機抽取履歷文字，再由顧問選擇保留原文或進行 AI 優化。",
        "支援中文 ATS 與英文 ATS 的新增人選／Add Profile，以及既有人選編輯頁面；候選人 URL 的 id 可變動，不需固定 id。",
        "支援一鍵 AI 優化、一鍵轉換語言，以及需要時還原原始履歷內容。",
        "可填入聯絡資料、摘要、技能、語言、工作經歷、教育經歷與明確存在的專案經歷。",
        "不自動按 ATS Save，不呼叫未授權的候選人建立 API；顧問確認後使用 ATS 原本的 Save／保存。",
    ]:
        add_bullet(doc, t)

    add_heading(doc, "3. 開始前準備")
    for t in [
        "使用最新版 Chrome，並確認已登入對應 ATS。",
        "安裝 Talent Nexus Connector 5000.0.113；更新後在 Chrome 擴充功能頁重新載入，再回到 ATS 重新整理一次頁面。",
        "確認目前位於新增／編輯人選頁，不是在候選人列表、登入頁或其他 ATS 功能頁。",
        "PDF／DOCX 檔案建議不超過 20 MB；掃描影像型 PDF 可能需要先做 OCR。",
    ]:
        add_bullet(doc, t)
    add_callout(doc, "語言選項提醒", "Connector 的介面語言不會限制 ATS 對應。請依目前開啟的中文或英文 ATS 頁面操作；履歷原文是中文就維持中文，英文就維持英文。", fill="FFF8E8")

    add_heading(doc, "4. 安裝與更新 Connector")
    for i, t in enumerate([
        "從公司 SharePoint 下載正式 ZIP，解壓縮到本機固定資料夾。",
        "在 Chrome 開啟 chrome://extensions，啟用右上角「開發人員模式」。",
        "選擇「載入解壓縮」，指定解壓縮後的 Connector 資料夾。",
        "若已安裝舊版本，按該擴充功能的重新載入按鈕，再重新整理 ATS 頁面。",
        "在擴充功能詳細資料確認版本顯示為 5000.0.113。",
    ], 1):
        add_number(doc, t)
    add_body(doc, "正式版本檔名：talent-nexus-connector-5000.0.113-ats-ai-fill.zip")

    add_heading(doc, "5. 新增人選：AI Fill 完整流程")
    add_heading(doc, "5.1 開啟面板並選取履歷", 2)
    for t in [
        "進入中文 ATS 的 candidate/edit 新增人選頁，或英文 ATS 的 Add Profile 頁。",
        "點擊上傳簡歷區域旁的「✨ AI Fill／✨ AI 填寫」按鈕；按鈕不會觸發 ATS 原始表單提交。",
        "拖放或選擇一份 PDF／DOCX。檔案選取前不會呼叫 AI。",
        "Connector 先在本機抽取文字並建立本次候選人 run ID；關閉面板或逾時後會清除暫存。",
    ]:
        add_number(doc, t)
    add_heading(doc, "5.2 選擇解析方式", 2)
    add_body(doc, "文字抽取完成後，請在面板選擇：")
    add_bullet(doc, "保留原始履歷：使用履歷原文與原始事實，不做語氣改寫。")
    add_bullet(doc, "AI 優化履歷：只在不改變真實性的前提下，把工作內容、資格經驗與專業優勢整理成較易閱讀的條列式文字。")
    add_body(doc, "若未選擇 AI 優化，Connector 仍可進行標準欄位解析；AI 不會在檔案剛選取時自動開始。")
    add_heading(doc, "5.3 預覽、優化與填入", 2)
    for i, t in enumerate([
        "在預覽區檢查姓名、電話、Email、地點、摘要、工作與教育紀錄。欄位狀態會標示已確認、需要選擇或履歷未提供。",
        "需要時使用「AI 優化」整理履歷、使用「轉換語言」轉為目標語言；可用「還原原始」回到上傳檔案的原始內容。",
        "按「填入 ATS」後，Connector 才將確認內容寫入目前頁面的表單；既有內容可依顧問選擇覆蓋。",
        "填入完成後由顧問逐欄確認，最後手動點擊 ATS 原本的 Save／保存。Connector 不會自動儲存。",
    ], 1):
        add_number(doc, t)

    add_heading(doc, "6. 既有人選編輯頁面")
    add_body(doc, "既有人選的補填方式與新增人選相同。中文 ATS 可從候選人列表進入目前可見的編輯表單；英文 ATS 可從 candidate/detail?id=任意有效 id 進入編輯表單。Connector 以目前頁面實際 Angular model 與表單結構辨識，不依賴固定 id。")
    add_callout(doc, "建議操作", "先確認候選人姓名與頁面上的既有內容，再按 AI Fill。若 ATS 原本解析的公司、日期或教育資訊有誤，可在預覽中允許覆蓋，填入後仍由顧問在 ATS 逐項確認。", fill="EEF6F8")

    add_heading(doc, "7. 欄位對應與資料來源")
    add_table(doc, ["履歷內容", "ATS 對應／處理原則"], [
        ("姓名、英文名、電話、Email、LinkedIn", "對應 candidate 基本資料；電話會移除連字號與空格，保留可驗證格式。"),
        ("現居地、期望地點、國外地址", "按履歷明確文字填入；國外地址保留英文，不翻成不確定的中文地名。"),
        ("摘要、技能、語言", "對應摘要／優勢與語言欄位；語言與熟練度若無可靠 ATS 選項，交由顧問選擇。"),
        ("工作經歷", "多筆建立 start、end、is_current、company_name、title、dept、description。"),
        ("教育經歷", "多筆建立 start、end、company_name（學校）、title（科系）；Degree 必須選 ATS 現有選項。"),
        ("專案經歷", "只有履歷明確寫出專案經歷時才填入；不從一般工作描述推測。"),
    ], [2700, 6660])
    add_body(doc, "Industry、Function、Location、Degree、Language、Proficiency 等分類欄位只能提出匹配建議。沒有可靠匹配時標記為「需要選擇」，不硬填任意 AI 文字。Source、Tags、Folder、Creator、Add to job 等顧問資料不由 AI Fill 覆寫。")

    add_heading(doc, "8. AI 優化與語言轉換規則")
    for t in [
        "優化只重新組織履歷中已存在的事實，不新增公司、職稱、技術、數字成果、證照或管理範圍。",
        "工作經歷優先條列化，將工作內容、資格經驗與專業優勢分開呈現，方便顧問快速檢查。",
        "中文履歷預設輸出中文，英文履歷預設輸出英文；姓名、公司名、產品名、證照名與技術名詞盡量保留原文。",
        "轉換語言只改變文字語言，不改變日期、公司、職稱、學校、電話、Email、LinkedIn 或其他鎖定事實。",
        "Email、電話、LinkedIn、公司、日期、學校、學位與地點等鎖定事實不得由 Gemini 生成或猜測。",
    ]:
        add_bullet(doc, t)

    add_heading(doc, "9. 電話、地點與來源識別")
    add_table(doc, ["項目", "處理方式"], [
        ("台灣手機", "例如 0913-755-058 會整理為 0913755058；+886 913 755 058 會轉為 0913755058。"),
        ("國外電話", "保留國碼與可驗證格式，不擅自改成台灣格式。"),
        ("國外地點", "履歷已使用英文時保留英文地址／城市；不依姓名或公司猜測所在地。"),
        ("104 來源", "底部保留 Connector 擷取的既有格式：`【104履歷代碼】<實際代碼>`；不由 Gemini 產生。"),
        ("LinkedIn 來源", "底部使用可驗證的 canonical profile URL：`【LinkedIn】https://www.linkedin.com/in/<slug>/`；無法驗證就省略。"),
    ], [2700, 6660])

    add_heading(doc, "10. 儲存、安全與資料邊界")
    for t in [
        "AI provider key 只存在後端，不放入 Connector、前端、localStorage 或 Git。",
        "PDF、抽取文字與解析狀態以候選人／run ID 隔離；關閉面板或逾時後清除暫存。",
        "不在 console、Netlify logs 或 analytics 記錄姓名、電話、Email、PDF 內容或完整履歷文字。",
        "AI 或 PDF 解析失敗時，不影響 ATS 原始表單與原始 Save；顧問可關閉面板後照常手動操作。",
        "不建立第二套未授權的 ATS 儲存 client，不直接呼叫候選人建立 API。",
    ]:
        add_bullet(doc, t)

    add_heading(doc, "11. 常見問題與排除方式")
    add_table(doc, ["現象", "處理方式"], [
        ("看不到 AI Fill", "確認位於 candidate/edit、Add Profile 或 candidate/detail 編輯頁；安裝／更新後重新載入 Connector，再重新整理 ATS。"),
        ("選檔後解析失敗", "先確認檔案未超過 20 MB；文字型 PDF／DOCX 可直接解析，掃描影像 PDF 需先 OCR 或改用可選取文字的版本。"),
        ("AI 優化失敗", "先保留原始履歷預覽，稍後重試；確認後端 API 可用。失敗不會改動 ATS 原始內容。"),
        ("下拉選項顯示需要選擇", "這是安全設計；由顧問在 ATS 現有選項中選擇，不要把 AI 自由文字硬塞入分類欄位。"),
        ("重跑後工作／教育重複", "正常流程會以目前 run 與既有紀錄比對；若仍重複，關閉面板、重新整理 ATS 後再做一次，並回報候選人頁面 URL。"),
        ("按填入後找不到資料", "確認按的是「填入 ATS」而非關閉；填入後請在 ATS 表單檢查，最後仍需手動 Save。"),
    ], [2700, 6660])

    add_heading(doc, "12. 顧問驗收清單")
    for t in [
        "頁面是正確候選人，姓名與來源一致。",
        "電話、Email、LinkedIn、現居地與期望地點與履歷一致。",
        "摘要、技能、語言及最高學歷沒有新增未提供的內容。",
        "每筆工作與教育紀錄的公司／學校、職稱／科系、日期與目前在職狀態正確。",
        "需要選擇的 Industry、Function、Location、Degree、Language、Proficiency 已由顧問確認。",
        "Source、Tags、Folder、Creator、Add to job 等顧問欄位沒有被覆寫。",
        "確認沒有自動 Save；完成檢查後再按 ATS 原本的 Save／保存。",
    ]:
        add_bullet(doc, t)
    add_callout(doc, "快速安全檢查", "如果 AI 服務暫時不可用，最安全的做法是關閉面板並繼續使用 ATS 原始表單；不要把未驗證的 AI 內容直接保存。", fill="FFF8E8")

    add_heading(doc, "13. 版本紀錄")
    add_table(doc, ["版本", "日期", "內容"], [
        ("v1.0", "2026-08-10", "首次發布完整操作手冊，涵蓋 ATS 新增／編輯 AI Fill、PDF／DOCX、AI 優化、語言轉換、來源識別與安全邊界。"),
    ], [1500, 1800, 6060])
    add_body(doc, "文件維護：Talent Nexus Connector 專案。若 Connector 版本、ATS 表單模型或 SharePoint 正式流程有重大變更，請建立新版本並保留本版歷史。")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
